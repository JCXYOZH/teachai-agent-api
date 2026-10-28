import io
from typing import Any, Dict, Tuple

import fitz
from docx import Document
from httpx import AsyncClient
from pptx import Presentation

from app.core.config import settings


async def register_user(
    client: AsyncClient,
    email: str = 'teacher@example.com',
    display_name: str = 'Teacher',
) -> Tuple[Dict[str, str], str]:
    response = await client.post(
        f'{settings.api_v1_prefix}/auth/register',
        json={
            'email': email,
            'password': 'Teach1234',
            'display_name': display_name,
        },
    )
    assert response.status_code == 201
    body = response.json()
    return {'Authorization': f"Bearer {body['access_token']}"}, body['user']['id']


def pdf_bytes(text: str) -> bytes:
    buffer = io.BytesIO()
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), text)
    document.save(buffer)
    document.close()
    return buffer.getvalue()


def docx_bytes(text: str) -> bytes:
    buffer = io.BytesIO()
    document = Document()
    document.add_paragraph(text)
    document.save(buffer)
    return buffer.getvalue()


def pptx_bytes(text: str) -> bytes:
    buffer = io.BytesIO()
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    box = slide.shapes.add_textbox(40, 40, 500, 80)
    box.text_frame.text = text
    presentation.save(buffer)
    return buffer.getvalue()


async def create_folder(client: AsyncClient, headers: Dict[str, str], name: str) -> str:
    response = await client.post(
        f'{settings.api_v1_prefix}/folders', headers=headers, json={'name': name}
    )
    assert response.status_code == 201
    return response.json()['id']


async def create_course(
    client: AsyncClient,
    headers: Dict[str, str],
    title: str,
    folder_id: str = None,
) -> str:
    payload: Dict[str, Any] = {'title': title}
    if folder_id is not None:
        payload['folder_id'] = folder_id
    response = await client.post(
        f'{settings.api_v1_prefix}/courses', headers=headers, json=payload
    )
    assert response.status_code == 201
    return response.json()['id']


async def upload_text(
    client: AsyncClient,
    headers: Dict[str, str],
    filename: str,
    content: str,
    course_id: str = None,
    folder_id: str = None,
) -> Dict[str, Any]:
    data: Dict[str, str] = {}
    if course_id:
        data['course_id'] = course_id
    if folder_id:
        data['folder_id'] = folder_id
    response = await client.post(
        f'{settings.api_v1_prefix}/assets',
        headers=headers,
        files={'file': (filename, content.encode('utf-8'), 'text/plain')},
        data=data,
    )
    assert response.status_code == 201, response.text
    return response.json()


async def test_folders_and_courses_are_owner_scoped(client: AsyncClient) -> None:
    headers, _ = await register_user(client)
    folder_id = await create_folder(client, headers, '数学讲义')

    duplicate = await client.post(
        f'{settings.api_v1_prefix}/folders', headers=headers, json={'name': '数学讲义'}
    )
    assert duplicate.status_code == 409
    assert duplicate.json()['error']['code'] == 'FOLDER_NAME_EXISTS'

    course_id = await create_course(client, headers, '极限与连续', folder_id)
    courses = await client.get(f'{settings.api_v1_prefix}/courses', headers=headers)
    assert courses.status_code == 200
    assert courses.json()['total'] == 1
    assert courses.json()['items'][0]['folder_name'] == '数学讲义'

    updated = await client.patch(
        f'{settings.api_v1_prefix}/courses/{course_id}',
        headers=headers,
        json={'title': '极限、连续与可导'},
    )
    assert updated.status_code == 200
    assert updated.json()['title'] == '极限、连续与可导'

    other_headers, _ = await register_user(client, 'other@example.com', 'Other')
    hidden = await client.get(
        f'{settings.api_v1_prefix}/courses/{course_id}', headers=other_headers
    )
    assert hidden.status_code == 404
    assert hidden.json()['error']['code'] == 'COURSE_NOT_FOUND'

    deleted = await client.delete(
        f'{settings.api_v1_prefix}/courses/{course_id}', headers=headers
    )
    assert deleted.status_code == 204
    remaining = await client.get(f'{settings.api_v1_prefix}/courses', headers=headers)
    assert remaining.json()['total'] == 0

    folder_deleted = await client.delete(
        f'{settings.api_v1_prefix}/folders/{folder_id}', headers=headers
    )
    assert folder_deleted.status_code == 204


async def test_upload_search_download_and_soft_delete(client: AsyncClient) -> None:
    headers, owner_id = await register_user(client)
    folder_id = await create_folder(client, headers, '函数资料')
    course_id = await create_course(client, headers, '函数复习', folder_id)

    asset = await upload_text(
        client,
        headers,
        '../../函数笔记.txt',
        '函数的极限是本节重点。连续函数在闭区间上有最大值。',
        course_id,
        folder_id,
    )
    assert asset['filename'] == '函数笔记.txt'
    assert asset['course_title'] == '函数复习'
    assert asset['folder_name'] == '函数资料'
    assert asset['extraction_status'] == 'ready'
    assert asset['char_count'] > 0

    expected_path = (
        settings.asset_storage_root
        / owner_id
        / asset['id']
        / 'original.txt'
    )
    assert expected_path.is_file()

    search = await client.get(
        f'{settings.api_v1_prefix}/search', headers=headers, params={'q': '极限'}
    )
    assert search.status_code == 200
    assert search.json()['total'] >= 1
    assert search.json()['assets'][0]['filename'] == '函数笔记.txt'
    assert '极限' in search.json()['assets'][0]['excerpt']

    detail = await client.get(
        f"{settings.api_v1_prefix}/assets/{asset['id']}", headers=headers
    )
    assert detail.status_code == 200
    assert '连续函数' in detail.json()['text_preview']
    assert detail.json()['sections'][0]['locator']['kind'] == 'lines'

    download = await client.get(
        f"{settings.api_v1_prefix}/assets/{asset['id']}/download", headers=headers
    )
    assert download.status_code == 200
    assert '极限' in download.content.decode('utf-8')
    assert 'attachment' in download.headers['content-disposition']

    other_headers, _ = await register_user(client, 'student@example.com', 'Student')
    hidden = await client.get(
        f"{settings.api_v1_prefix}/assets/{asset['id']}", headers=other_headers
    )
    assert hidden.status_code == 404
    hidden_download = await client.get(
        f"{settings.api_v1_prefix}/assets/{asset['id']}/download",
        headers=other_headers,
    )
    assert hidden_download.status_code == 404

    deleted = await client.delete(
        f"{settings.api_v1_prefix}/assets/{asset['id']}", headers=headers
    )
    assert deleted.status_code == 204
    assert expected_path.is_file()
    assert (await client.get(
        f"{settings.api_v1_prefix}/assets/{asset['id']}", headers=headers
    )).status_code == 404


async def test_folder_delete_unfiles_courses_and_assets(client: AsyncClient) -> None:
    headers, _ = await register_user(client)
    folder_id = await create_folder(client, headers, '临时目录')
    course_id = await create_course(client, headers, '待整理课程', folder_id)
    asset = await upload_text(
        client, headers, 'outline.txt', '课程提纲', course_id, folder_id
    )

    response = await client.delete(
        f'{settings.api_v1_prefix}/folders/{folder_id}', headers=headers
    )
    assert response.status_code == 204

    course = await client.get(
        f'{settings.api_v1_prefix}/courses/{course_id}', headers=headers
    )
    assert course.json()['folder_id'] is None
    asset_detail = await client.get(
        f"{settings.api_v1_prefix}/assets/{asset['id']}", headers=headers
    )
    assert asset_detail.json()['folder_id'] is None


async def test_pdf_docx_and_pptx_are_parsed(client: AsyncClient) -> None:
    headers, _ = await register_user(client)
    cases = [
        ('calculus.pdf', pdf_bytes('Calculus limit theorem'), 'pdf', 1),
        ('lesson.docx', docx_bytes('中文教案：导数与变化率'), 'docx', None),
        ('slides.pptx', pptx_bytes('课堂演示：函数图像'), 'pptx', 1),
    ]

    for filename, payload, extension, page_count in cases:
        response = await client.post(
            f'{settings.api_v1_prefix}/assets',
            headers=headers,
            files={'file': (filename, payload, 'application/octet-stream')},
        )
        assert response.status_code == 201, response.text
        body = response.json()
        assert body['extension'] == f'.{extension}'
        assert body['extraction_status'] == 'ready'
        assert body['page_count'] == page_count

    search = await client.get(
        f'{settings.api_v1_prefix}/search', headers=headers, params={'q': '导数'}
    )
    assert search.status_code == 200
    assert any(item['filename'] == 'lesson.docx' for item in search.json()['assets'])


async def test_utf16_text_is_extracted(client: AsyncClient) -> None:
    headers, _ = await register_user(client)
    response = await client.post(
        f'{settings.api_v1_prefix}/assets',
        headers=headers,
        files={
            'file': (
                'utf16-notes.txt',
                'UTF-16 中文材料：矩阵与线性变换。'.encode('utf-16'),
                'text/plain',
            )
        },
    )
    assert response.status_code == 201, response.text
    asset = response.json()
    assert asset['extraction_status'] == 'ready'

    detail = await client.get(
        f"{settings.api_v1_prefix}/assets/{asset['id']}", headers=headers
    )
    assert detail.status_code == 200
    assert '线性变换' in detail.json()['text_preview']


async def test_upload_rejects_unsupported_and_mismatched_files(client: AsyncClient) -> None:
    headers, owner_id = await register_user(client)

    unsupported = await client.post(
        f'{settings.api_v1_prefix}/assets',
        headers=headers,
        files={'file': ('script.exe', b'MZ', 'application/octet-stream')},
    )
    assert unsupported.status_code == 415
    assert unsupported.json()['error']['code'] == 'UNSUPPORTED_FILE_TYPE'

    mismatched = await client.post(
        f'{settings.api_v1_prefix}/assets',
        headers=headers,
        files={'file': ('fake.pdf', b'not a pdf', 'application/pdf')},
    )
    assert mismatched.status_code == 422
    assert mismatched.json()['error']['code'] == 'FILE_PARSE_FAILED'
    owner_dir = settings.asset_storage_root / owner_id
    assert not owner_dir.exists() or not any(owner_dir.iterdir())


async def test_upload_enforces_size_limit(
    client: AsyncClient, monkeypatch: Any
) -> None:
    headers, _ = await register_user(client)
    with monkeypatch.context() as patch:
        patch.setattr(settings, 'asset_max_bytes', 4)
        response = await client.post(
            f'{settings.api_v1_prefix}/assets',
            headers=headers,
            files={'file': ('too-large.txt', b'12345', 'text/plain')},
        )
    assert response.status_code == 413
    assert response.json()['error']['code'] == 'FILE_TOO_LARGE'
