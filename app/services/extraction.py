import codecs
import re
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import fitz
from docx import Document
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


class ExtractionError(Exception):
    pass


@dataclass(frozen=True)
class ExtractedDocument:
    text: str
    sections: List[Dict[str, Any]]
    page_count: Optional[int]
    truncated: bool


def _decode_text(data: bytes) -> str:
    encodings: Tuple[str, ...]
    if data.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)):
        encodings = ('utf-16',)
    elif 0 in data:
        raise ExtractionError('The text file contains binary data')
    else:
        encodings = ('utf-8-sig', 'gb18030')

    for encoding in encodings:
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ExtractionError('The text file encoding is not supported')


def extract_document(
    path: Path,
    extension: str,
    max_chars: int,
) -> ExtractedDocument:
    if extension == ".pdf":
        raw_sections, page_count = _extract_pdf(path)
    elif extension == ".docx":
        raw_sections, page_count = _extract_docx(path)
    elif extension == ".pptx":
        raw_sections, page_count = _extract_pptx(path)
    elif extension in {".md", ".markdown", ".txt"}:
        raw_sections, page_count = _extract_text(path)
    else:
        raise ExtractionError("Unsupported file type")
    return _limit_sections(raw_sections, page_count, max_chars)


def _check_pdf(path: Path) -> None:
    with path.open("rb") as stream:
        if stream.read(5) != b"%PDF-":
            raise ExtractionError("The file content is not a valid PDF")


def _check_office_archive(path: Path, extension: str) -> None:
    try:
        with zipfile.ZipFile(path) as archive:
            names = set(archive.namelist())
            required = (
                "word/document.xml"
                if extension == ".docx"
                else "ppt/presentation.xml"
            )
            if "[Content_Types].xml" not in names or required not in names:
                raise ExtractionError("The file content does not match its extension")
            total_uncompressed = sum(info.file_size for info in archive.infolist())
            if total_uncompressed > 500 * 1024 * 1024:
                raise ExtractionError("The document expands beyond the safety limit")
            if any(info.flag_bits & 0x1 for info in archive.infolist()):
                raise ExtractionError("Encrypted Office documents are not supported")
    except zipfile.BadZipFile as exc:
        raise ExtractionError("The file content is not a valid Office document") from exc


def _extract_pdf(path: Path) -> Tuple[List[Tuple[Dict[str, Any], str]], Optional[int]]:
    _check_pdf(path)
    sections: List[Tuple[Dict[str, Any], str]] = []
    try:
        with fitz.open(str(path)) as document:
            if document.needs_pass:
                raise ExtractionError("Password-protected PDFs are not supported")
            for index, page in enumerate(document, start=1):
                text = page.get_text("text").strip()
                if text:
                    sections.append(
                        (
                            {
                                "kind": "page",
                                "index": index,
                                "label": f"第 {index} 页",
                            },
                            text,
                        )
                    )
            return sections, document.page_count
    except ExtractionError:
        raise
    except Exception as exc:
        raise ExtractionError("The PDF could not be parsed") from exc


def _extract_docx(path: Path) -> Tuple[List[Tuple[Dict[str, Any], str]], Optional[int]]:
    _check_office_archive(path, ".docx")
    try:
        document = Document(str(path))
    except Exception as exc:
        raise ExtractionError("The DOCX file could not be parsed") from exc

    blocks: List[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            blocks.append(text)
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            text = " | ".join(cell for cell in cells if cell)
            if text:
                blocks.append(text)
    text = "\n".join(blocks)
    sections = (
        [({"kind": "document", "index": 1, "label": "全文"}, text)] if text else []
    )
    return sections, None


def _extract_pptx(path: Path) -> Tuple[List[Tuple[Dict[str, Any], str]], Optional[int]]:
    _check_office_archive(path, ".pptx")
    try:
        presentation = Presentation(str(path))
    except Exception as exc:
        raise ExtractionError("The PPTX file could not be parsed") from exc

    sections: List[Tuple[Dict[str, Any], str]] = []
    for index, slide in enumerate(presentation.slides, start=1):
        blocks: List[str] = []
        _collect_shape_text(slide.shapes, blocks)
        try:
            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    blocks.append(notes)
        except (AttributeError, KeyError, ValueError):
            pass
        text = "\n".join(block for block in blocks if block.strip())
        if text:
            sections.append(
                (
                    {
                        "kind": "slide",
                        "index": index,
                        "label": f"第 {index} 页幻灯片",
                    },
                    text,
                )
            )
    return sections, len(presentation.slides)


def _collect_shape_text(shapes: Any, blocks: List[str]) -> None:
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            _collect_shape_text(shape.shapes, blocks)
            continue
        text = getattr(shape, "text", "")
        if text and text.strip():
            blocks.append(text.strip())
        if getattr(shape, "has_table", False):
            for row in shape.table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                row_text = " | ".join(cell for cell in cells if cell)
                if row_text:
                    blocks.append(row_text)


def _extract_text(path: Path) -> Tuple[List[Tuple[Dict[str, Any], str]], Optional[int]]:
    data = path.read_bytes()
    text = _decode_text(data)

    sections: List[Tuple[Dict[str, Any], str]] = []
    lines = text.splitlines()
    for start in range(0, len(lines), 80):
        chunk = "\n".join(lines[start : start + 80]).strip()
        if not chunk:
            continue
        end = min(start + 80, len(lines))
        sections.append(
            (
                {
                    "kind": "lines",
                    "start": start + 1,
                    "end": end,
                    "label": f"第 {start + 1}-{end} 行",
                },
                chunk,
            )
        )
    return sections, None


def _limit_sections(
    raw_sections: List[Tuple[Dict[str, Any], str]],
    page_count: Optional[int],
    max_chars: int,
) -> ExtractedDocument:
    sections: List[Dict[str, Any]] = []
    remaining = max_chars
    truncated = False
    for locator, raw_text in raw_sections:
        text = re.sub(r"\n{3,}", "\n\n", raw_text).strip()
        if not text:
            continue
        if remaining <= 0:
            truncated = True
            break
        clipped = text[:remaining]
        if len(clipped) < len(text):
            truncated = True
        sections.append(
            {
                "locator": locator,
                "text": clipped,
                "char_count": len(clipped),
            }
        )
        remaining -= len(clipped)
        if truncated:
            break

    full_text = "\n\n".join(section["text"] for section in sections)
    return ExtractedDocument(
        text=full_text,
        sections=sections,
        page_count=page_count,
        truncated=truncated,
    )
